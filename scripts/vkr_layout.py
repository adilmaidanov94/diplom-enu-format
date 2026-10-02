#!/usr/bin/env python3
"""Постраничная вёрстка дипломной работы после vkr_format.py.

python3 vkr_layout.py formatted.docx final.docx [--kind work|project] [--max-iter 40]
        [--report layout.json] [--pdf final.pdf] [+ любые флаги vkr_format.py через --fmt "..."]

Цикл: рендер в PDF (LibreOffice) → поиск первой проблемной страницы → исправление → снова рендер.
Проблемы: страница начинается с рисунка / таблицы / формулы / «где» / подписи;
пустота внизу страницы > 3 строк, когда следующая страница не начинается с нового раздела;
таблица перешла на новую страницу без «Продолжение таблицы N» (каз.: «N-кестенің жалғасы»).
"""
import argparse, copy, json, os, re, shlex, shutil, subprocess, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import vkr_format as F

MM = 72 / 25.4
TOP = 20 * MM
BOTTOM = 841.89 - 25 * MM
LINE = 16.1
VOID_LINES = 3

def word_like_copy(docx, outdir):
    """LibreOffice переносит абзац с «не отрывать от следующего» ЦЕЛИКОМ, а Word — только последние строки
    (с запретом висячих строк — 2). Для честной имитации Word в технической копии последние ~2 строки
    такого абзаца выделяются в отдельный абзац. Сам документ не меняется."""
    doc = Document(docx)
    for p in list(doc.element.body.iterchildren(qn('w:p'))):
        kn = p.find(qn('w:pPr') + '/' + qn('w:keepNext'))
        if kn is None or kn.get(qn('w:val')) in ('0', 'false'):
            continue
        txt = F.ptext(p)
        if len(txt) < 4 * F.CHARS_PER_LINE or F.X(p, './/w:drawing|.//m:oMath|.//w:fldChar'):
            continue
        cut = txt.rfind(' ', 0, len(txt) - int(1.6 * F.CHARS_PER_LINE))
        if cut <= 0:
            continue
        runs = F._split_runs_at(p, [cut + 1])
        tail = F.new_par()
        pPr = p.find(qn('w:pPr'))
        if pPr is not None:
            tail.append(copy.deepcopy(pPr))
            F.set_ppr(tail, 'ind', {'left': 0, 'firstLine': 0, 'right': 0})
        pos = 0
        for r in runs:
            rt = ''.join(F.X(r, './w:t/text()'))
            if pos >= cut + 1:
                tail.append(r)
            pos += len(rt)
        F.flag(p, 'keepNext', False)
        F.set_ppr(p, 'jc', {'val': 'both'})
        p.addnext(tail)
    sim = os.path.join(outdir, 'sim.docx')
    doc.save(sim)
    return sim

def render(docx, outdir):
    docx = word_like_copy(docx, outdir)
    prof = tempfile.mkdtemp(prefix='lo_')
    subprocess.run(['soffice', f'-env:UserInstallation=file://{prof}', '--headless', '--convert-to', 'pdf',
                    '--outdir', outdir, docx], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    shutil.rmtree(prof, ignore_errors=True)
    return os.path.join(outdir, os.path.splitext(os.path.basename(docx))[0] + '.pdf')

MATH_FONTS = ('opensymbol', 'cambria math', 'stixmath', 'symbol')

def page_model(page):
    lines = []
    d = page.get_text('dict')
    for b in d['blocks']:
        if b['type'] != 0:
            continue
        for l in b['lines']:
            txt = ''.join(s['text'] for s in l['spans']).strip()
            if not txt:
                continue
            x0, y0, x1, y1 = l['bbox']
            if y0 > BOTTOM + 4 or y1 < TOP - 8:
                continue        # колонтитулы, номер страницы
            fonts = {s['font'].lower() for s in l['spans']}
            lines.append(dict(y0=y0, y1=y1, x0=x0, x1=x1, text=txt, math=any(any(m in f for m in MATH_FONTS) for f in fonts)))
    # объединить фрагменты одной строки
    lines.sort(key=lambda l: (round(l['y0']), l['x0']))
    merged = []
    for l in lines:
        if merged and abs(merged[-1]['y0'] - l['y0']) < 3:
            m = merged[-1]; m['text'] += ' ' + l['text']; m['x1'] = max(m['x1'], l['x1']); m['y1'] = max(m['y1'], l['y1'])
            m['math'] = m['math'] or l['math']
        else:
            merged.append(dict(l))
    graphics = []
    for im in page.get_image_info():
        x0, y0, x1, y1 = im['bbox']
        if y1 - y0 > 8 and y0 < BOTTOM + 4:
            graphics.append(dict(y0=y0, y1=y1, x0=x0, x1=x1, kind='image'))
    for dr in page.get_drawings():
        r = dr['rect']
        if r.y0 > BOTTOM + 4 or r.y1 < TOP - 8:
            continue
        if r.width < 3 and r.height < 3:
            continue
        graphics.append(dict(y0=r.y0, y1=r.y1, x0=r.x0, x1=r.x1, kind='draw', w=r.width, h=r.height))
    return merged, graphics

def is_grid(graphics, y_from):
    """Есть ли вверху табличная сетка (много тонких горизонтальных/вертикальных линий)."""
    g = [x for x in graphics if x['kind'] == 'draw' and x['y0'] < y_from + 40]
    thin = [x for x in g if x['h'] < 2.5 or x['w'] < 2.5]
    return len(thin) >= 3

def classify(pages):
    out = []
    for pi, (lines, graphics) in enumerate(pages):
        info = dict(page=pi + 1, first='', start=None, bottom=TOP, n_lines=len(lines))
        tops = []
        if lines:
            tops.append(('text', lines[0]['y0']))
        if graphics:
            gtop = min(g['y0'] for g in graphics)
            tops.append(('graphic', gtop))
        if not tops:
            info['start'] = 'blank'; out.append(info); continue
        bottom = max([l['y1'] for l in lines] + [g['y1'] for g in graphics])
        info['bottom'] = bottom
        info['first'] = lines[0]['text'] if lines else ''
        first_kind = min(tops, key=lambda t: t[1])
        if first_kind[0] == 'graphic' and (not lines or first_kind[1] < lines[0]['y0'] - 1):
            gtop = first_kind[1]
            if is_grid(graphics, gtop) and lines and any(g['kind'] == 'draw' for g in graphics if g['y0'] < lines[0]['y0']):
                info['start'] = 'table'
            else:
                info['start'] = 'figure'
        else:
            t = lines[0]['text']
            if F.CONT_RE.match(t):
                info['start'] = 'cont'
            elif F.TBLCAP_RE.match(t) and len(t) < 250:
                info['start'] = 'tblcap'
            elif F.FIGCAP_RE.match(t) and len(t) < 250 and (len(lines) < 2 or not re.match(r'^[а-яё]', lines[1]['text'])):
                info['start'] = 'figcap'
            elif F.WHERE_RE.match(t) and not re.match(r'^где-', t, re.I):
                info['start'] = 'where'
            elif lines[0]['math'] or (re.search(r'\(\d+(\.\d+)?\)\s*$', t) and len(t) < 90 and (lines[0]['x0'] > 150)):
                info['start'] = 'formula'
            else:
                info['start'] = 'text'
                # сколько строк текста до первой графики
                gs = [g['y0'] for g in graphics if g['y0'] > lines[0]['y0'] and (g['kind'] == 'image' or g.get('h', 0) > 8)]
                if gs:
                    gy = min(gs)
                    info['lines_before_graphic'] = sum(1 for l in lines if l['y1'] <= gy + 1)
        info['lines'] = [l['text'] for l in lines[:6]]
        info['all_lines'] = lines
        out.append(info)
    return out

HEAD1 = re.compile(r'^(\d{1,2}\s+\S|' + F.STRUCT_RE.pattern[1:-1] + r'|приложение\b)', re.I)

def find_problem(cls, heading_pages, intro_page, skip=()):
    """Первая проблема сверху вниз (страницы до «Введения» не трогаем)."""
    for i, c in enumerate(cls):
        if c['page'] < intro_page:
            continue
        s = c['start']
        sig = lambda t: (t, c['first'][:40], c['n_lines'] and c['all_lines'][-1]['text'][:40])
        if s in ('figure', 'table', 'tblcap', 'figcap', 'where', 'formula') and c['page'] not in heading_pages \
                and sig('starts_with_' + s) not in skip:
            return dict(page=c['page'], type='starts_with_' + s, sig=sig('starts_with_' + s))
        if c.get('lines_before_graphic') == 1 and sig('one_line') not in skip:
            return dict(page=c['page'], type='one_line_before_graphic', sig=sig('one_line'))
        if i + 1 < len(cls) and cls[i + 1]['page'] not in heading_pages and cls[i + 1]['start'] != 'blank':
            void = BOTTOM - c['bottom']
            if void > VOID_LINES * LINE and sig('void') not in skip:
                return dict(page=c['page'], type='void', void=void, sig=sig('void'))
    return None

def heading_pages_of(cls, doc_headings):
    """Страницы, начинающиеся с заголовка 1-го уровня (там пустота на предыдущей странице допустима)."""
    pages = set()
    for c in cls:
        if not c.get('lines'):
            continue
        f = F.norm(c['lines'][0]).lower()
        for h in doc_headings:
            if f.startswith(h.lower()[:25]):
                pages.add(c['page']); break
    return pages

def doc_structure(path):
    doc = Document(path)
    items, intro = F.analyze(doc)
    blocks = F.build_blocks(items)
    for b in blocks:
        b['key'] = F.block_key(b, items)
    h1 = [it.text for it in items if it.kind in ('heading', 'appx') and it.level == 1]
    return doc, items, blocks, h1

def block_on_page(c, blocks, items):
    """Какая вставка на странице c (первая), по подписи."""
    keys = []
    for l in c.get('all_lines', []):
        t = l['text']
        m = F.FIGCAP_RE.match(t)
        if m and len(t) < 250:
            keys.append(('Рисунок ' + m.group(2), l['y0']))
        m = F.TBLCAP_RE.match(t)
        if m and len(t) < 250:
            keys.append(('Таблица ' + m.group(2), l['y0']))
        m = re.search(r'\((\d+(?:\.\d+)?)\)\s*$', t)
        if m and (l['math'] or l['x0'] > 150) and len(t) < 90:
            keys.append(('Формула ' + m.group(1), l['y0']))
    known = {b['key'] for b in blocks if b['key']}
    keys = [k for k in keys if k[0] in known]
    return keys

def scale_block_image(doc, items, blocks, key, factor):
    for b in blocks:
        if b['key'] == key and b['type'] == 'fig':
            done = False
            for x in b['fig']:
                for inl in F.X(items[x].el, './/wp:inline'):
                    F.scale_inline(inl, factor); done = True
            return done
    return False

TCPR_ORDER = ['cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar',
              'textDirection', 'tcFitText', 'vAlign', 'hideMark']

def _tcpr_child(tcPr, name):
    el = tcPr.find(qn('w:' + name))
    if el is None:
        el = OxmlElement('w:' + name)
        idx = TCPR_ORDER.index(name); pos = 0
        for i, ch in enumerate(tcPr):
            n = F.tag(ch)
            if n in TCPR_ORDER and TCPR_ORDER.index(n) < idx:
                pos = i + 1
        tcPr.insert(pos, el)
    return el

def split_table(doc, items, blocks, num, first_texts):
    """Разрыв таблицы num перед строкой, с которой началась новая страница."""
    parts = [b for b in blocks if b['type'] == 'tbl' and b['key'] in ('Таблица ' + num, 'Продолжение Таблица ' + num)]
    cand = [F.norm(t).lower() for t in first_texts if F.norm(t)]
    def row_matches(cells, line):
        toks = line.split()
        if not toks or toks[0] != cells[0].split()[0]:
            return False
        pos = 0
        for c in cells[:5]:
            k = line.find(c[:min(8, len(c))], pos)
            if k < 0:
                return False
            pos = k + 1
        return True
    for b in reversed(parts):
        tbl = items[b['tbl']].el
        rows = F.X(tbl, './w:tr')
        for ri, tr in enumerate(rows):
            if ri == 0:
                continue
            cells = [F.norm(F.ptext(tc)).lower() for tc in F.X(tr, './w:tc')]
            cells = [c for c in cells if c]
            if not cells:
                continue
            if any(row_matches(cells, t) for t in cand):
                new = copy.deepcopy(tbl)
                for r in F.X(new, './w:tr')[:ri]:
                    new.remove(r)
                # повторить шапку
                if ri > 1 and not b.get('cont'):
                    hdr = copy.deepcopy(rows[0])
                    F.X(new, './w:tr')[0].addprevious(hdr)
                elif b.get('cont'):
                    hdr = copy.deepcopy(rows[0])
                    F.X(new, './w:tr')[0].addprevious(hdr)
                for r in rows[ri:]:
                    tbl.remove(r)
                # вертикальные объединения в первой перенесённой строке
                body_rows = F.X(new, './w:tr')
                first_data = body_rows[1] if len(body_rows) > 1 and (ri > 1 or b.get('cont')) else body_rows[0]
                for vm in F.X(first_data, './w:tc/w:tcPr/w:vMerge'):
                    vm.set(qn('w:val'), 'restart')
                # нижнюю черту первой части не проводят
                for tc in F.X(F.X(tbl, './w:tr')[-1], './w:tc'):
                    tcPr = tc.find(qn('w:tcPr'))
                    if tcPr is None:
                        tcPr = OxmlElement('w:tcPr'); tc.insert(0, tcPr)
                    brd = _tcpr_child(tcPr, 'tcBorders')
                    bt = brd.find(qn('w:bottom'))
                    if bt is None:
                        bt = OxmlElement('w:bottom'); brd.append(bt)
                    bt.set(qn('w:val'), 'nil')
                cap = F.new_par(F.cont_caption(num, F.doc_lang(doc)))
                tbl.addnext(cap)
                cap.addnext(new)
                return True
    return False

def run(args):
    work = tempfile.mkdtemp(prefix='vkr_')
    cur = os.path.join(work, 'cur.docx')
    shutil.copy(args.inp, cur)
    d0 = Document(cur)
    if F.merge_continuations(d0):
        d0.save(cur)
    fmt_extra = shlex.split(args.fmt or '')
    split_ok, moved, scaled, log = set(), {}, {}, []
    seen = {}
    skip = set()
    final_cls = None
    for it in range(args.max_iter):
        pdf = render(cur, work)
        pdoc = pymupdf.open(pdf)
        pages = [page_model(p) for p in pdoc]
        cls = classify(pages)
        doc, items, blocks, h1 = doc_structure(cur)
        hp = heading_pages_of(cls, h1)
        intro_page = min([c['page'] for c in cls if c.get('lines') and F.INTRO_RE.match(F.norm(c['lines'][0]))] or [1])
        prob = find_problem(cls, hp, intro_page, skip)
        final_cls = (cls, hp, intro_page, h1)
        if prob is None:
            break
        sig = (prob['page'], prob['type'])
        seen[sig] = seen.get(sig, 0) + 1
        if seen[sig] > 4:
            log.append(dict(prob, action='не удалось исправить автоматически'))
            skip.add(prob['sig']); continue
        c = cls[prob['page'] - 1]
        action = None
        if prob['type'] == 'void' and prob['page'] < len(cls) and cls[prob['page']]['start'] == 'cont':
            num = F.CONT_RE.match(cls[prob['page']]['first']).group(2)
            if F.merge_continuations(doc, num):
                action = f'таблица {num}: разрыв устарел — таблица склеена для нового разрыва'
                log.append(dict(prob, action=action)); doc.save(cur)
                subprocess.run([sys.executable, os.path.join(os.path.dirname(__file__), 'vkr_format.py'), cur, cur,
                                '--split-ok', ';'.join(split_ok)] + fmt_extra, check=True, stdout=subprocess.DEVNULL)
                continue
        if prob['type'] == 'starts_with_table':
            # таблица перешла на новую страницу без «Продолжение таблицы»
            num = None
            for pc in reversed(cls[:prob['page'] - 1]):
                for l in reversed(pc.get('all_lines', [])):
                    m = F.TBLCAP_RE.match(l['text']) or F.CONT_RE.match(l['text'])
                    if m:
                        num = m.group(2); break
                if num:
                    break
            if num:
                lines = c['all_lines']
                top = [l['text'] for l in lines if l['y0'] < lines[0]['y0'] + 3] + [l['text'] for l in lines[:3]]
                if split_table(doc, items, blocks, num, top):
                    action = f'таблица {num}: разрыв с подписью «{F.cont_caption(num, F.doc_lang(doc))}»'
                    split_ok.add('Таблица ' + num)
        elif prob['type'] in ('void', 'starts_with_figure', 'starts_with_figcap', 'starts_with_tblcap',
                              'starts_with_formula', 'starts_with_where', 'one_line_before_graphic'):
            tgt_page = prob['page'] + 1 if prob['type'] == 'void' else prob['page']
            keys = block_on_page(cls[tgt_page - 1], blocks, items) if tgt_page <= len(cls) else []
            if prob['type'] != 'void' and prob['page'] > 1 and not keys:
                keys = block_on_page(cls[prob['page'] - 2], blocks, items)[-1:]
            if keys:
                key = keys[0][0]
                b = next(b for b in blocks if b['key'] == key)
                nxt = F.block_end(b) + 1
                can_move = b['type'] == 'fig' and nxt < len(items) and items[nxt].kind == 'text' and not F.CONT_RE.match(items[nxt].text) \
                    and moved.get(key, 0) < 3 and items[nxt].zone == items[F.block_start(b)].zone
                if prob['type'] == 'void' and b['type'] == 'tbl' and key not in split_ok and 3 < len(F.X(items[b['tbl']].el, './w:tr')) <= F.SMALL_TABLE_ROWS:
                    split_ok.add(key); action = f'{key}: разрешён разрыв таблицы (пустота {prob["void"]/LINE:.0f} строк)'
                elif prob['type'] == 'void' and can_move:
                    moved[key] = moved.get(key, 0) + 1
                    fmt_extra_once = ['--move-after', key]
                    action = f'{key}: перенесена на абзац ниже (пустота {prob["void"]/LINE:.0f} строк на стр. {prob["page"]})'
                    doc.save(cur)
                    subprocess.run([sys.executable, os.path.join(os.path.dirname(__file__), 'vkr_format.py'), cur, cur,
                                    '--split-ok', ';'.join(split_ok)] + fmt_extra + fmt_extra_once,
                                   check=True, stdout=subprocess.DEVNULL)
                    log.append(dict(prob, action=action))
                    continue
                elif b['type'] == 'fig' and scaled.get(key, 1.0) > 0.62:
                    k = 0.85
                    if prob['type'] == 'void':
                        imgs = [g for g in pages[tgt_page - 1][1] if g['kind'] == 'image']
                        if imgs:
                            h = imgs[0]['y1'] - imgs[0]['y0']
                            need = (imgs[0]['y1'] - TOP) + 3 * LINE - prob['void']
                            k = max(0.6 / scaled.get(key, 1.0), min(0.95, (h - need) / h)) if h > need else 0.85
                    if scale_block_image(doc, items, blocks, key, k):
                        scaled[key] = scaled.get(key, 1.0) * k
                        action = f'{key}: уменьшен до {scaled[key]*100:.0f}%'
        if action is None:
            log.append(dict(prob, action='не удалось исправить автоматически'))
            skip.add(prob['sig']); continue
        log.append(dict(prob, action=action))
        doc.save(cur)
        subprocess.run([sys.executable, os.path.join(os.path.dirname(__file__), 'vkr_format.py'), cur, cur,
                        '--split-ok', ';'.join(split_ok)] + fmt_extra, check=True, stdout=subprocess.DEVNULL)
    # финальный отчёт
    pdf = render(cur, work)
    pdoc = pymupdf.open(pdf)
    pages = [page_model(p) for p in pdoc]
    cls = classify(pages)
    doc, items, blocks, h1 = doc_structure(cur)
    hp = heading_pages_of(cls, h1)
    intro_page = min([c['page'] for c in cls if c.get('lines') and F.INTRO_RE.match(F.norm(c['lines'][0]))] or [1])
    remaining, minor = [], []
    for i, c in enumerate(cls):
        if c['page'] < intro_page:
            continue
        if c['start'] in ('figure', 'table', 'tblcap', 'figcap', 'where', 'formula') and c['page'] not in hp:
            remaining.append(f'стр. {c["page"]}: начинается с «{c["start"]}» — {c["first"][:60]}')
        if c.get('lines_before_graphic') == 1:
            remaining.append(f'стр. {c["page"]}: перед вставкой только 1 строка текста')
        if i + 1 < len(cls) and cls[i + 1]['page'] not in hp and cls[i + 1]['start'] != 'blank':
            void = BOTTOM - c['bottom']
            if void > VOID_LINES * LINE:
                nxt_start = cls[i + 1]['start']
                if void <= 6.5 * LINE:
                    minor.append(f'стр. {c["page"]}: пустота ≈{void/LINE:.0f} строк — неразрывный блок «текст + вставка» не помещается')
                else:
                    remaining.append(f'стр. {c["page"]}: пустота внизу ≈{void/LINE:.0f} строк')
    # объём по страницам
    starts = []
    for h in h1:
        for c in cls:
            if c.get('lines') and F.norm(c['lines'][0]).lower().startswith(h.lower()[:25]):
                starts.append((h, c['page'])); break
    vol = {}
    for k, (h, p) in enumerate(starts):
        end = starts[k + 1][1] - 1 if k + 1 < len(starts) else len(cls)
        vol[h] = end - p + 1
    intro_p = next((p for h, p in starts if F.INTRO_RE.match(h)), None)
    refs_p = next((p for h, p in starts if F.REFS_RE.match(h)), None)
    total = ((refs_p or len(cls) + 1) - intro_p) if intro_p else None
    vol_issues = []
    minimum = 50 if args.kind == 'work' else 40
    if total is not None and total < minimum:
        vol_issues.append(f'Объём от введения до заключения включительно: {total} стр. — нужно не менее {minimum} (п.9.14)')
    for h, n in vol.items():
        if F.INTRO_RE.match(h) and not 1 <= n <= 3:
            vol_issues.append(f'Введение: {n} стр. — рекомендуется 1–3 (п.9.9)')
        if re.match(r'^(заключение|қорытынды|conclusion)$', h, re.I) and not 1 <= n <= 3:
            vol_issues.append(f'Заключение: {n} стр. — рекомендуется 1–3')
    shutil.copy(cur, args.out)
    if args.pdf:
        shutil.copy(pdf, args.pdf)
    for a_ in log: a_.pop('sig', None)
    rep = dict(pages=len(cls), iterations=len(log), actions=log, remaining=remaining, minor_voids=minor, volume_pages=vol,
               main_pages=total, volume_issues=vol_issues)
    js = json.dumps(rep, ensure_ascii=False, indent=1)
    if args.report:
        open(args.report, 'w', encoding='utf-8').write(js)
    print(js)

if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('inp'); ap.add_argument('out')
    ap.add_argument('--kind', default='work', choices=['work', 'project'])
    ap.add_argument('--max-iter', type=int, default=40)
    ap.add_argument('--report'); ap.add_argument('--pdf')
    ap.add_argument('--fmt', default='', help='доп. флаги для vkr_format.py, например "--start-page 9"')
    run(ap.parse_args())
